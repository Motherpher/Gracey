from __future__ import annotations

import json
import os
import re
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt

from .hrdm import public_candidate_evidence


def _safe_name(value: str) -> str:
    value = re.sub(r"[^\w\- ]+", "", value or "job").strip()
    return re.sub(r"\s+", "_", value)[:70] or "job"


def run_ai_application(job: dict, profile: dict, hrdm: dict, lane: str) -> dict | None:
    if not os.getenv("OPENAI_API_KEY"):
        return None
    try:
        from openai import OpenAI
    except Exception:
        return None

    model = os.getenv("CAREERHUB_MODEL") or "gpt-5.4"
    client = OpenAI()
    schema = {
        "type": "object",
        "additionalProperties": False,
        "required": ["cover_letter", "cv_profile", "cv_bullets", "interview_notes", "claims_check"],
        "properties": {
            "cover_letter": {"type": "string"},
            "cv_profile": {"type": "string"},
            "cv_bullets": {"type": "array", "items": {"type": "string"}},
            "interview_notes": {"type": "array", "items": {"type": "string"}},
            "claims_check": {"type": "array", "items": {"type": "string"}},
        },
    }
    length_instruction = "Håll brevet kompakt och praktiskt." if lane == "bridge" else "Skriv ett fokuserat professionellt brev på högst en sida."
    verified_profile = public_candidate_evidence(profile)
    prompt = f"""Skapa ett ansökningsunderlag utifrån den genomförda matchningsanalysen.

Regler:
- Använd endast verifierad karriärevidens i profilen nedan.
- Hitta aldrig på datum, arbetsgivare, verktyg, språknivå, utbildning, kvalifikationer eller resultat.
- Bevara osäkerhet som osäkerhet.
- Search-only önskemål eller behov är aldrig kandidatfakta och får inte bli ansökningspåståenden.
- {length_instruction}
- Skriv naturlig, vuxen och idiomatisk svenska med lugn självsäkerhet.
- Prioritera belägg som är direkt relevanta för rollen.
- Nämn inte HRDM i ansökningsbrevet.
- Returnera endast strukturerad JSON.

JOBB:
{json.dumps(job, ensure_ascii=False, indent=2)}

VERIFIERAD KARRIÄRPROFIL:
{json.dumps(verified_profile, ensure_ascii=False, indent=2)}

HRDM:
{json.dumps(hrdm, ensure_ascii=False, indent=2)}
"""
    response = client.responses.create(
        model=model,
        store=False,
        input=prompt,
        text={"format":{"type":"json_schema","name":"careerhub_application","strict":True,"schema":schema}},
    )
    return json.loads(response.output_text)


def fallback_application(job: dict, profile: dict, hrdm: dict) -> dict:
    name = profile.get("identity", {}).get("name", "Grace")
    role = job.get("title") or "the role"
    company = job.get("company") or "your organisation"
    matches = hrdm.get("candidate_positioning", {}).get("strong_matches", [])
    evidence = "\n".join(f"- {x}" for x in matches[:4]) or "- Add evidence after HRDM analysis."
    return {
        "cover_letter": (
            f"Ansökan – {role} hos {company}\n\n"
            f"Det här är ett tillfälligt utkast för {name}. Den automatiska ansökningstexten kördes inte. "
            "När analysen är klar ersätts texten med ett underlag som bygger på verifierad karriärevidens.\n\n"
            f"Belägg att utgå från:\n{evidence}"
        ),
        "cv_profile": "Profiltext förbereds när den verifierade analysen är klar.",
        "cv_bullets": [],
        "interview_notes": [],
        "claims_check": ["Använd endast verifierade karriäruppgifter i den slutliga ansökan."],
    }


def _base_doc(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(title, level=0)
    return doc


def write_application_docx(outdir: Path, job: dict, profile: dict, app: dict) -> Path:
    role = job.get("title") or "Role"
    company = job.get("company") or "Employer"
    doc = _base_doc(f"Ansökningsutkast — {role}")
    p = doc.add_paragraph()
    p.add_run(company).bold = True
    if job.get("url"):
        doc.add_paragraph(job["url"])

    doc.add_heading("Personligt brev – utkast", level=1)
    for para in app.get("cover_letter", "").split("\n\n"):
        doc.add_paragraph(para)

    doc.add_heading("CV-profil – utkast", level=1)
    doc.add_paragraph(app.get("cv_profile", ""))

    if app.get("cv_bullets"):
        doc.add_heading("CV – lyft fram", level=1)
        for x in app["cv_bullets"]:
            doc.add_paragraph(x, style="List Bullet")

    if app.get("claims_check"):
        doc.add_heading("Kontroll före skick", level=1)
        for x in app["claims_check"]:
            doc.add_paragraph(x, style="List Bullet")

    path = outdir / f"Application_{_safe_name(company)}_{_safe_name(role)}.docx"
    doc.save(path)
    return path


def write_hrdm_docx(outdir: Path, hrdm: dict) -> Path:
    job = hrdm.get("job", {})
    doc = _base_doc(f"HRDM-R brief — {job.get('title') or 'Role'}")
    sections = [
        ("Hidden need", hrdm.get("hidden_need", {})),
        ("FunctionCore", hrdm.get("function_core", {})),
        ("DoD", hrdm.get("dod", {})),
        ("Assessment zones", hrdm.get("assessment_zones", [])),
        ("Grace i relation till rollen", hrdm.get("candidate_positioning", {})),
        ("HCC", hrdm.get("hcc", {})),
        ("Ansökningsstrategi", hrdm.get("application_strategy", {})),
    ]
    for heading, value in sections:
        doc.add_heading(heading, level=1)
        if isinstance(value, dict):
            for k, v in value.items():
                p = doc.add_paragraph()
                p.add_run(str(k).replace("_", " ").title() + ": ").bold = True
                p.add_run(json.dumps(v, ensure_ascii=False, indent=2) if isinstance(v, (list, dict)) else str(v))
        elif isinstance(value, list):
            for x in value:
                doc.add_paragraph(json.dumps(x, ensure_ascii=False) if isinstance(x, dict) else str(x), style="List Bullet")
        else:
            doc.add_paragraph(str(value))

    path = outdir / f"HRDM_{_safe_name(job.get('company',''))}_{_safe_name(job.get('title','Role'))}.docx"
    doc.save(path)
    return path
